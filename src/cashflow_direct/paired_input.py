"""把摘要式现流明细的原项目接回原始序时账现金分录。"""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from cashflow_direct.money import yuan_to_cent
from cashflow_direct.normalization import CASH_ACCOUNT_TERMS


FLOW_FIELDS = {
    "voucher": ("凭证号", "凭证编号", "凭证字号"),
    "item": ("现金流量项目", "现金流量表项目", "现流项目"),
    "summary": ("摘要", "业务摘要", "业务说明"),
    "direction": ("方向", "收付方向", "流入流出方向"),
    "amount": ("金额", "现流金额", "现金流量金额"),
}
JOURNAL_FIELDS = {
    "date": ("会计月", "日期", "记账日期", "凭证日期"),
    "voucher": ("凭证编号", "凭证号", "凭证号码"),
    "word": ("凭证种类", "凭证类别", "凭证字"),
    "summary": ("业务说明", "摘要", "业务摘要"),
    "account": ("科目名称", "会计科目名称", "完整科目名称"),
    "code": ("科目编号", "科目编码", "科目代码"),
    "debit": ("借方发生额", "借方金额", "借方"),
    "credit": ("贷方发生额", "贷方金额", "贷方"),
}
PARENT_NAMES = ("上级科目名称", "科目路径", "上级科目")
ROOT_NAMES = ("一级科目", "一级科目名称")
DATE_NAMES = ("日期", "记账日期", "凭证日期", "会计月")


@dataclass(frozen=True)
class PairResult:
    workbook: Path
    report: Path
    journal_rows: int
    matched_cash_rows: int
    flow_rows: int
    duplicate_groups: int


def _headers(row: tuple[object, ...]) -> dict[str, int]:
    return {str(value).strip().replace(" ", ""): i for i, value in enumerate(row) if value is not None}


def _field(headers: dict[str, int], names: tuple[str, ...], required: bool = True) -> int | None:
    found = [headers[name] for name in names if name in headers]
    if len(set(found)) > 1:
        raise ValueError(f"字段有多个候选列：{'、'.join(names)}；请先整理或明确表头")
    if not found and required:
        raise ValueError(f"缺少字段：{'／'.join(names)}")
    return found[0] if found else None


def _value(row: tuple[object, ...], column: int | None) -> object:
    return row[column] if column is not None and column < len(row) else None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _day(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    match = re.match(r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if not match:
        raise ValueError(f"日期无法识别：{text!r}")
    return date(*(int(part) for part in match.groups()))


def _voucher(value: object, word: object = None) -> tuple[str, str]:
    raw = _text(value)
    kind = _text(word)
    match = re.match(r"^([^\d]*?)[-－—\s]*(\d+)$", raw)
    if match:
        kind = kind or match.group(1).strip("-－— ")
        number = str(int(match.group(2)))
    elif raw.isdigit():
        number = str(int(raw))
    else:
        raise ValueError(f"凭证号无法识别：{raw!r}")
    return kind, number


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _flow_header(headers: dict[str, int]) -> tuple[int | None, int | None, int | None]:
    one = _field(headers, DATE_NAMES, False)
    month = next((index for name, index in headers.items() if re.fullmatch(r"\d{4}年月", name)), None)
    day = next((index for name, index in headers.items() if re.fullmatch(r"\d{4}年日", name)), None)
    if one is None and (month is None or day is None):
        raise ValueError("现流明细缺少可识别的日期或年份月／年份日列")
    return one, month, day


def _flow_date(row: tuple[object, ...], headers: dict[str, int], columns: tuple[int | None, int | None, int | None]) -> date:
    one, month, day = columns
    if one is not None:
        return _day(_value(row, one))
    assert month is not None and day is not None
    year_month = re.match(r"(\d{4})年月", next(name for name, index in headers.items() if index == month))
    year_day = re.match(r"(\d{4})年日", next(name for name, index in headers.items() if index == day))
    if year_month is None or year_day is None or year_month.group(1) != year_day.group(1):
        raise ValueError("年月列与年日列的年份不一致")
    return date(int(year_month.group(1)), int(_value(row, month)), int(_value(row, day)))


def _sheet(path: Path, fields: dict[str, tuple[str, ...]], selected: str | None, raw_journal: bool = False):
    book = load_workbook(path, read_only=True, data_only=True)
    if selected:
        if selected not in book.sheetnames:
            book.close()
            raise ValueError(f"{path.name} 中没有工作表《{selected}》")
        names = [selected]
    else:
        names = book.sheetnames
    candidates = []
    for name in names:
        sheet = book[name]
        for header_row, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), 1):
            headers = _headers(row)
            if all(any(alias in headers for alias in aliases) for aliases in fields.values()):
                if raw_journal and any(label in headers for label in ("对方科目", "对方科目名称")):
                    break
                candidates.append((sheet, headers, header_row))
                break
    if len(candidates) != 1:
        book.close()
        raise ValueError(f"{path.name} 可用的{'原始序时账' if raw_journal else '现流明细'}工作表有 {len(candidates)} 张；请指定工作表，且原始账不得使用生成的对方科目拆分表")
    return book, *candidates[0]


def _account_path(row: tuple[object, ...], headers: dict[str, int], leaf_col: int) -> str:
    leaf = _text(_value(row, leaf_col)).strip("\\/_")
    parent_col = _field(headers, PARENT_NAMES, False)
    root_col = _field(headers, ROOT_NAMES, False)
    parent = _text(_value(row, parent_col)).strip("\\/_")
    root = _text(_value(row, root_col)).strip("\\/_")
    parts = [part for part in re.split(r"[\\/]", parent) if part]
    if not parts and root:
        parts = [root]
    if leaf and (not parts or parts[-1] != leaf):
        parts.append(leaf)
    return "_".join(parts)


def prepare_paired_input(
    flow_path: Path,
    journal_path: Path,
    output_path: Path,
    *,
    flow_sheet: str | None = None,
    journal_sheet: str | None = None,
) -> PairResult:
    """逐现金分录核对原项目，保留原始分录，不构造业务拆分。"""
    flow_path, journal_path, output_path = map(Path, (flow_path, journal_path, output_path))
    if output_path.exists() or output_path.with_suffix(".json").exists():
        raise FileExistsError("配对输出已存在，请选择新的输出文件名；不会覆盖已有文件")
    if flow_path.resolve() == journal_path.resolve():
        raise ValueError("现流明细与原始序时账必须分别指定")
    flow_hash, journal_hash = _sha256(flow_path), _sha256(journal_path)
    flow_book, flow_ws, fh, flow_header_row = _sheet(flow_path, FLOW_FIELDS, flow_sheet)
    journal_book = None
    try:
        journal_book, journal_ws, jh, journal_header_row = _sheet(journal_path, JOURNAL_FIELDS, journal_sheet, True)
        fcols = {role: _field(fh, names) for role, names in FLOW_FIELDS.items()}
        jcols = {role: _field(jh, names) for role, names in JOURNAL_FIELDS.items()}
        date_cols = _flow_header(fh)
        flow_groups: dict[tuple[object, ...], list[tuple[int, str]]] = defaultdict(list)
        flow_dates: list[date] = []
        for number, row in enumerate(flow_ws.iter_rows(min_row=flow_header_row + 1, values_only=True), flow_header_row + 1):
            if not _text(_value(row, fcols["voucher"])):
                continue  # 无凭证号的小计行不是业务明细
            day = _flow_date(row, fh, date_cols)
            kind, voucher = _voucher(_value(row, fcols["voucher"]))
            direction = _text(_value(row, fcols["direction"]))
            if direction not in {"流入", "流出"}:
                raise ValueError(f"现流明细第 {number} 行方向无法识别：{direction!r}")
            amount = yuan_to_cent(_value(row, fcols["amount"]))
            if amount == 0:
                raise ValueError(f"现流明细第 {number} 行金额不能为零")
            if amount < 0:
                direction = "流出" if direction == "流入" else "流入"
                amount = -amount
            key = (day, kind, voucher, _text(_value(row, fcols["summary"])), direction, amount)
            flow_groups[key].append((number, _text(_value(row, fcols["item"]))))
            flow_dates.append(day)
        if not flow_dates:
            raise ValueError("现流明细没有可配对的凭证行")
        start = min(flow_dates).replace(day=1)
        last = max(flow_dates)
        end = date(last.year + (last.month == 12), last.month % 12 + 1, 1)
        journal_rows: list[tuple[int, tuple[object, ...], tuple[object, ...] | None, str]] = []
        cash_groups: dict[tuple[object, ...], list[int]] = defaultdict(list)
        for number, row in enumerate(journal_ws.iter_rows(min_row=journal_header_row + 1, values_only=True), journal_header_row + 1):
            if not _text(_value(row, jcols["voucher"])):
                continue
            day = _day(_value(row, jcols["date"]))
            if not start <= day < end:
                continue
            kind, voucher = _voucher(_value(row, jcols["voucher"]), _value(row, jcols["word"]))
            path = _account_path(row, jh, jcols["account"])
            key = None
            if any(path.split("_")[0] == term for term in CASH_ACCOUNT_TERMS):
                debit = yuan_to_cent(_value(row, jcols["debit"]) or 0)
                credit = yuan_to_cent(_value(row, jcols["credit"]) or 0)
                if debit and credit:
                    raise ValueError(f"原始序时账第 {number} 行现金分录借贷金额异常")
                if debit or credit:
                    net = debit - credit
                    key = (day, kind, voucher, _text(_value(row, jcols["summary"])), "流入" if net > 0 else "流出", abs(net))
                    cash_groups[key].append(number)
            journal_rows.append((number, row, key, path))
        for key, rows in flow_groups.items():
            if len({item for _, item in rows}) > 1:
                raise ValueError(f"同日同凭证同摘要同方向同金额对应多个原项目，不能确定原始现金分录项目：{key[0]} {key[1]}-{key[2]}")
        flow_lines = [(key, source, item) for key, rows in flow_groups.items() for source, item in rows]
        cash_lines = [(key, source) for key, rows in cash_groups.items() for source in rows]
        used_flow: set[int] = set()
        assigned: dict[int, tuple[str, str, str]] = {}

        def match_equal(group_key, method: str) -> int:
            flow_buckets: dict[tuple[object, ...], list[tuple[tuple[object, ...], int, str]]] = defaultdict(list)
            cash_buckets: dict[tuple[object, ...], list[tuple[tuple[object, ...], int]]] = defaultdict(list)
            for key, source, item in flow_lines:
                if source not in used_flow:
                    flow_buckets[group_key(key)].append((key, source, item))
            for key, source in cash_lines:
                if source not in assigned:
                    cash_buckets[group_key(key)].append((key, source))
            count = 0
            for grouped, flows in flow_buckets.items():
                cash = cash_buckets.get(grouped, ())
                if len(flows) != len(cash):
                    continue
                items = {item for _, _, item in flows}
                if len(items) != 1:
                    raise ValueError(f"同一配对组对应多个原项目，不能确定原始现金分录项目：{grouped}")
                source_rows = ",".join(str(source) for _, source, _ in flows)
                for _, source in cash:
                    assigned[source] = (next(iter(items)), source_rows, method)
                used_flow.update(source for _, source, _ in flows)
                count += len(flows)
            return count

        exact_count = match_equal(lambda key: key, "逐行")
        month_count = match_equal(
            lambda key: ((key[0].year, key[0].month), *key[1:]), "同月凭证"
        )

        def voucher_month(key: tuple[object, ...]) -> tuple[object, ...]:
            return (key[0].year, key[0].month, key[1], key[2])

        residual_flow: dict[tuple[object, ...], list[tuple[tuple[object, ...], int, str]]] = defaultdict(list)
        residual_cash: dict[tuple[object, ...], list[tuple[tuple[object, ...], int]]] = defaultdict(list)
        for key, source, item in flow_lines:
            if source not in used_flow:
                residual_flow[voucher_month(key)].append((key, source, item))
        for key, source in cash_lines:
            if source not in assigned:
                residual_cash[voucher_month(key)].append((key, source))
        net_groups = 0
        unresolved: list[str] = []
        for voucher, flows in residual_flow.items():
            cash = residual_cash.get(voucher, ())
            flow_net = sum(key[5] if key[4] == "流入" else -key[5] for key, _, _ in flows)
            cash_net = sum(key[5] if key[4] == "流入" else -key[5] for key, _ in cash)
            if not cash or flow_net != cash_net:
                unresolved.append(f"{voucher[0]}-{voucher[1]:02d} {voucher[2]}-{voucher[3]}：现流净额{flow_net / 100:.2f}，原账剩余现金净额{cash_net / 100:.2f}")
                continue
            by_direction: dict[str, list[tuple[int, str]]] = defaultdict(list)
            for key, source, item in flows:
                by_direction[str(key[4])].append((source, item))
            for direction, rows in by_direction.items():
                if len({item for _, item in rows}) > 1:
                    raise ValueError(f"同凭证同方向存在多个原项目，不能把净额对应到原始现金行：{voucher} {direction}")
                if not any(key[4] == direction for key, _ in cash):
                    raise ValueError(f"现流方向在原始现金分录中没有对应现金行：{voucher} {direction}")
            for direction, rows in by_direction.items():
                item = rows[0][1]
                source_rows = ",".join(str(source) for source, _ in rows)
                for key, source in cash:
                    if key[4] == direction:
                        assigned[source] = (item, source_rows, "凭证净额")
            used_flow.update(source for _, source, _ in flows)
            net_groups += 1
        if unresolved:
            raise ValueError("现流与原始现金分录无法核对：" + "；".join(unresolved[:5]))
        if len(used_flow) != len(flow_lines):
            raise ValueError("仍有现流明细未能关联原始现金分录")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        out = Workbook()
        sheet = out.active
        sheet.title = "配对后原始序时账"
        sheet.append(("日期", "凭证字", "凭证号", "摘要", "科目编码", "科目名称", "借方发生额", "贷方发生额", "现金流量项目", "原始序时账行号", "现流明细匹配组行号", "原项目关联方式"))
        for number, row, key, path in journal_rows:
            day = _day(_value(row, jcols["date"]))
            kind, voucher = _voucher(_value(row, jcols["voucher"]), _value(row, jcols["word"]))
            item, sources, method = assigned.get(number, ("", "", ""))
            sheet.append((day.isoformat(), kind, voucher, _text(_value(row, jcols["summary"])), _text(_value(row, jcols["code"])), path, _value(row, jcols["debit"]) or 0, _value(row, jcols["credit"]) or 0, item, number, sources, method))
        out.save(output_path)
        out.close()
        if _sha256(flow_path) != flow_hash or _sha256(journal_path) != journal_hash:
            output_path.unlink()
            raise ValueError("输入文件在配对过程中发生变化，请重新执行")
        report = output_path.with_suffix(".json")
        audit = {
            "flow_file": str(flow_path.resolve()), "flow_sha256": flow_hash, "flow_sheet": flow_ws.title,
            "journal_file": str(journal_path.resolve()), "journal_sha256": journal_hash, "journal_sheet": journal_ws.title,
            "period_start": start.isoformat(), "period_end_exclusive": end.isoformat(),
            "flow_rows": len(flow_dates), "journal_rows": len(journal_rows),
            "matched_cash_rows": len(assigned), "exact_matched_flow_rows": exact_count,
            "month_matched_flow_rows": month_count, "net_matched_vouchers": net_groups,
            "duplicate_groups": sum(len(rows) > 1 for rows in flow_groups.values()),
            "unlabeled_cash_rows": len(cash_lines) - len(assigned),
            "pairing": "先核对日期、凭证、摘要、有效方向和整数分金额；再核对同月凭证的相同明细；仍不逐行相等时，仅在凭证净额相等且每方向原项目唯一时关联现金行，不拆分原始分录",
        }
        report.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
        return PairResult(output_path, report, len(journal_rows), len(assigned), len(flow_dates), audit["duplicate_groups"])
    finally:
        flow_book.close()
        if journal_book is not None:
            journal_book.close()
