from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from cashflow_direct.pipeline import (
    confirm_cash_scope,
    finalize_run,
    run_classification,
    run_preflight,
)
from scripts.串行上下文调度 import initialize
from tests.fixture_factory import mark_dictionary_complete, write_ai_end_to_end_case
from tests.test_pipeline import _complete_all_pending_ai


def _prepare_ai_run(root: Path, source: Path):
    # 两条路径必须读取同一份输入文件：file_id 含输入文件绝对路径（intake.py），
    # 输入路径不同会派生出不同的业务编号，那是输入差异而不是 1M/128K 差异。
    root.mkdir(parents=True, exist_ok=True)
    preflight = run_preflight(
        [source], ("1000000", "750000", "50000"), output_parent=root
    )
    confirm_cash_scope(preflight.run_dir, preflight.recommended_cash_decisions)
    mark_dictionary_complete(preflight.run_dir)
    run_classification(preflight.run_dir)
    return preflight


def business_decisions(run_dir: Path) -> dict:
    state = json.loads(
        (run_dir / "计算留痕数据" / "运行状态.json").read_text(encoding="utf-8-sig")
    )
    return {
        "decisions": state.get("decisions"),
        "reconciliation": state.get("reconciliation"),
        "classification_summary": state.get("classification_summary"),
    }


def workbook_business_signature(path: Path, root: Path) -> list:
    workbook = load_workbook(path, data_only=False)
    metadata_labels = {"运行编号", "生成时间", "结果路径", "文件位置"}

    def normalize(value):
        return value.replace(str(root), "<ROOT>") if isinstance(value, str) else value

    signature = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            if row[0].value in metadata_labels:
                continue
            rows.append([
                (
                    normalize(cell.value),
                    cell.data_type,
                    cell.number_format,
                    str(cell.font),
                    str(cell.fill),
                    str(cell.border),
                    str(cell.alignment),
                    str(cell.protection),
                    normalize(cell.hyperlink.target) if cell.hyperlink else None,
                )
                for cell in row
            ])
        validations = [
            (item.type, normalize(item.formula1), str(item.sqref))
            for item in sheet.data_validations.dataValidation
        ]
        signature.append({
            "title": sheet.title,
            "state": sheet.sheet_state,
            "freeze": str(sheet.freeze_panes),
            "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
            "rows": rows,
            "validations": validations,
        })
    workbook.close()
    if "03 现金流量表正表" not in [item["title"] for item in signature]:
        raise AssertionError("未实际读取03 现金流量表正表，不能用空签名通过")
    return signature


def test_same_inputs_and_fixed_ai_results_have_same_business_output(tmp_path: Path):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    source = write_ai_end_to_end_case(input_dir)
    baseline_root = tmp_path / "baseline"
    optimized_root = tmp_path / "optimized"
    baseline_run = _prepare_ai_run(baseline_root, source)
    _complete_all_pending_ai(baseline_run.run_dir, baseline_root)
    baseline_final = finalize_run(baseline_run.run_dir)

    optimized_run = _prepare_ai_run(optimized_root, source)
    initialize(optimized_run.run_dir, "128K", "test-current-model", optimized_root / "backup")
    _complete_all_pending_ai(optimized_run.run_dir, optimized_root)
    optimized_final = finalize_run(optimized_run.run_dir)

    baseline_results = {p.name: p for p in baseline_root.glob("自动完成*.jsonl")}
    optimized_results = {p.name: p for p in optimized_root.glob("自动完成*.jsonl")}
    assert baseline_results
    assert baseline_results.keys() == optimized_results.keys()
    for name, baseline_file in baseline_results.items():
        assert baseline_file.read_bytes() == optimized_results[name].read_bytes()

    assert workbook_business_signature(
        baseline_final.workbook_path, baseline_root
    ) == workbook_business_signature(optimized_final.workbook_path, optimized_root)
    assert business_decisions(baseline_run.run_dir) == business_decisions(optimized_run.run_dir)


def test_workbook_signature_detects_injected_difference(tmp_path: Path):
    root = tmp_path / "single"
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    source = write_ai_end_to_end_case(input_dir)
    preflight = _prepare_ai_run(root, source)
    _complete_all_pending_ai(preflight.run_dir, root)
    final = finalize_run(preflight.run_dir)

    plain_copy = tmp_path / "plain_copy.xlsx"
    shutil.copy2(final.workbook_path, plain_copy)
    assert workbook_business_signature(final.workbook_path, root) == workbook_business_signature(plain_copy, root)

    round_trip = tmp_path / "round_trip.xlsx"
    workbook = load_workbook(final.workbook_path, data_only=False)
    workbook.save(round_trip)
    workbook.close()
    assert workbook_business_signature(final.workbook_path, root) == workbook_business_signature(round_trip, root)

    mutated = tmp_path / "mutated.xlsx"
    shutil.copy2(round_trip, mutated)
    workbook = load_workbook(mutated, data_only=False)
    sheet = workbook["03 现金流量表正表"]
    changed = False
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, (int, float)) and cell.value:
                cell.value = cell.value + 1
                changed = True
                break
        if changed:
            break
    assert changed, "03 现金流量表正表中没有可修改的数值单元格"
    workbook.save(mutated)
    workbook.close()
    assert workbook_business_signature(final.workbook_path, root) != workbook_business_signature(mutated, root)
