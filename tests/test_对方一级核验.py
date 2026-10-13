"""对方一级核验只使用已确认的现金分配，不用同凭证无关业务凑匹配。"""
import json

import pytest
from openpyxl import Workbook, load_workbook

from cashflow_direct.components import build_cashflow_components, confirm_cash_scope, discover_cash_scope
from cashflow_direct.pipeline import run_preflight, confirm_cash_scope as confirm_run_cash, supplement_cash_balances, run_classification, finalize_run
from tests.fixture_factory import _component_entry


def entry(n, account, amount, summary, counterpart=""):
    return _component_entry(n, "V", account, debit_cent=max(amount, 0), credit_cent=max(-amount, 0),
                            counterpart_name=counterpart, summary=summary)


def build(rows):
    proposal = discover_cash_scope(rows)
    scope = confirm_cash_scope(proposal, {c.account_key: "include" for c in proposal.candidates})
    return build_cashflow_components(rows, scope)


def checks(result):
    records = getattr(result, "counterpart_level1_checks", None)
    assert records is not None, "尚未实现对方一级核验"
    return {c.cash_entry_id: c for c in records}


@pytest.mark.parametrize("hint,expected", [("管理费用", "一致"), ("固定资产", "不一致"), ("管理费用_差旅费", "一致"), ("", "未提供")])
def test_只比较一级科目并保留来源(hint, expected):
    rows = (entry(1, "1002 银行存款_活期", -10000, "支付办公费", hint),
            entry(2, "管理费用_办公费", 10000, "支付办公费"))
    result = build(rows)
    check = checks(result)["E1"]
    assert check.status == expected
    assert check.actual_level1s == ("管理费用",)
    assert check.source == rows[0].source
    assert check.declared_counterpart == hint
    component = result.components[0]
    assert ("counterpart_level1_mismatch" in component.anomalies) == (expected == "不一致")
    assert component.cash_delta_cent == -10000
    assert component.counterpart_accounts == ("管理费用_办公费",)


def test_销售收入与销项税集合不误报矛盾():
    rows = (entry(1, "1002 银行存款_活期", 11300, "收销售货款", "主营业务收入"),
            entry(2, "主营业务收入_商品收入", -10000, "收销售货款"),
            entry(3, "应交税费_应交增值税_销项税额", -1300, "收销售货款"))
    result = build(rows)
    check = checks(result)["E1"]
    assert check.status == "一致"
    assert set(check.actual_level1s) == {"主营业务收入", "应交税费"}
    assert all("counterpart_level1_mismatch" not in c.anomalies for c in result.components)


def test_一条现金来源跨多个业务按集合核验():
    rows = (entry(1, "1002 银行存款_活期", -10000, "支付办公用品及设备", "管理费用"),
            entry(2, "管理费用_办公费", 6000, "支付办公用品及设备"),
            entry(3, "固定资产_办公设备", 4000, "支付办公用品及设备"))
    result = build(rows)
    assert len(result.components) == 2
    check = checks(result)["E1"]
    assert check.status == "一致"
    assert set(check.actual_level1s) == {"管理费用", "固定资产"}
    assert len(check.component_ids) == 2
    assert all("counterpart_level1_mismatch" not in c.anomalies for c in result.components)


def test_不能拿同凭证另一笔业务来掩盖矛盾():
    rows = (entry(1, "1002 银行存款_A", -10000, "支付办公费", "固定资产"),
            entry(2, "管理费用_办公费", 10000, "支付办公费"),
            entry(3, "1002 银行存款_B", -5000, "支付设备款", "固定资产"),
            entry(4, "固定资产_办公设备", 5000, "支付设备款"))
    result = build(rows)
    records = checks(result)
    assert records["E1"].status == "不一致"
    assert records["E1"].actual_level1s == ("管理费用",)
    assert records["E3"].status == "一致"
    by_summary = {c.summary: c for c in result.components}
    assert "counterpart_level1_mismatch" in by_summary["支付办公费"].anomalies
    assert "counterpart_level1_mismatch" not in by_summary["支付设备款"].anomalies


def test_对应关系未确定时不能仅按同凭证科目判错():
    rows = (entry(1, "1002 银行存款_A", -8000, "支付款项", "固定资产"),
            entry(2, "1002 银行存款_B", -7000, "支付款项", "固定资产"),
            entry(3, "应付账款_货款", 10000, "支付款项"),
            entry(4, "应付职工薪酬_工资", 5000, "支付款项"))
    result = build(rows)
    assert result.structure_requests
    assert {c.status for c in checks(result).values()} == {"待核实"}
    assert all("counterpart_level1_mismatch" not in c.anomalies for c in result.components)


def test_只有现金行不能把填入的对方一级当成已验证业务路径():
    result = build((entry(1, "1002 银行存款_活期", -10000, "支付办公费", "管理费用"),))
    check = checks(result)["E1"]
    assert check.status == "待核实"
    assert check.actual_level1s == ()


def test_核验不受输入行顺序影响():
    rows = (entry(1, "1002 银行存款_活期", -10000, "支付办公费", "固定资产"),
            entry(2, "管理费用_办公费", 10000, "支付办公费"))
    assert checks(build(rows)) == checks(build(tuple(reversed(rows))))


def test_核验结果经过正式入口保存并提供中文异常说明(tmp_path):
    source = tmp_path / "合成凭证.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "序时账"
    sheet.append(["日期", "凭证号", "科目名称", "科目编码", "摘要", "借方", "贷方", "对方科目", "现金流量项目"])
    sheet.append(["2025-01-01", "记-1", "银行存款_活期", "100201", "支付办公费", 0, 100, "固定资产", "支付其他与经营活动有关的现金"])
    sheet.append(["2025-01-01", "记-1", "管理费用_办公费", "660201", "支付办公费", 100, 0, None, None])
    workbook.save(source)
    result = run_preflight([source], (250000, 125000, 10000), output_parent=tmp_path)
    assert result.status == "waiting_cash_scope"
    confirm_run_cash(result.run_dir, {k: "include" for k in result.recommended_cash_decisions})
    supplement_cash_balances(result.run_dir, 100, 0, 0, "合成测试的独立余额")
    run_classification(result.run_dir)
    state = json.loads((result.run_dir / "计算留痕数据" / "运行状态.json").read_text(encoding="utf-8-sig"))
    records = state.get("counterpart_level1_checks", [])
    assert len(records) == 1
    assert records[0]["status"] == "不一致"
    trace = result.run_dir / "计算留痕数据" / "对方科目一级核验.jsonl"
    assert trace.exists()
    assert json.loads(trace.read_text(encoding="utf-8-sig").strip())["declared_level1"] == "固定资产"
    final = finalize_run(result.run_dir)
    output = load_workbook(final.workbook_path, read_only=True, data_only=True)
    try:
        rows = output["全量分类留痕"].iter_rows(values_only=True)
        header = next(rows)
        row = next(rows)
        values = dict(zip(header, row))
        assert "对方一级科目与对应业务不一致" in values["异常"]
        assert "counterpart_level1_mismatch" not in values["异常"]
        assert values["原现流项目"] == "支付其他与经营活动有关的现金"
    finally:
        output.close()
