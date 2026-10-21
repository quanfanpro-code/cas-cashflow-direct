"""匿名回归：断流接续、逐来源守恒、结构待办与诊断导出。"""


from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from cashflow_direct import pipeline as p
from tests.fixture_factory import mark_dictionary_complete
from tests.test_pipeline import _complete_all_pending_ai

def prepare(tmp_path):
    w=Workbook();s=w.active;s.append(['日期','凭证号','摘要','科目名称','借方金额','贷方金额'])
    for v,summary,account,debit,credit in [('待查','支付费用','银行存款',0,70),('待查','支付费用','银行存款',0,30),('待查','经营推广','管理费用_推广费',70,0),('待查','单位代垫款','其他应收款_代垫款',30,0),('已知','支付银行手续费','财务费用_手续费',5,0),('已知','支付银行手续费','银行存款',0,5)]:
        s.append(['2026-01-01',v,summary,account,debit,credit])
    f=tmp_path/'匿名序时账.xlsx';w.save(f);w.close()
    run=p.run_preflight([f],('250000','125000','10000'),tmp_path).run_dir
    state=p._load_state(run);p.confirm_cash_scope(run,{c['account_key']:'include' for c in state['cash_scope_proposal']['candidates']})
    p.supplement_cash_balances(run,'200','95','0','匿名测试余额');mark_dictionary_complete(run);p.run_classification(run)
    return run

def test_诊断保留结构门禁与现金(tmp_path):
    run=prepare(tmp_path);_complete_all_pending_ai(run,tmp_path)
    before=p._load_state(run);assert before['component_structure_requests'];assert not before.get('component_structure_ai_tasks')
    result=p.export_diagnostic_run(run);after=p._load_state(run)
    assert after['stage']==before['stage'] and after['component_structure_requests']==before['component_structure_requests']
    assert 'workbook_path' not in after and after['diagnostic_workbook_path']==str(result.workbook_path)
    assert result.overall_status.startswith('诊断材料')
    assert sum(c['cash_delta_cent'] for c in after['components'])==-10500
    w=load_workbook(result.workbook_path,read_only=True,data_only=False)
    assert len(w.sheetnames)==14 and str(w['使用说明与状态']['B3'].value).startswith('诊断材料');w.close()
    with pytest.raises(RuntimeError,match='业务组成'):p.finalize_run(run)

def test_诊断不得绕过分类AI(tmp_path):
    run=prepare(tmp_path);s=p._load_state(run)
    s['classification_summary']['ai_tasks_missing']=1;p._save_state(run,s)
    with pytest.raises(RuntimeError,match='AI'):p.export_diagnostic_run(run)

def test_正式入口仍拒绝未确认结构(tmp_path):
    run=prepare(tmp_path)
    with pytest.raises(RuntimeError,match='业务组成'):p.finalize_run(run)

from dataclasses import replace
from tests.test_blank_original_fallback import decision,component,ORDERED_ITEMS,DIRECTIONS
from cashflow_direct.blank_original_fallback import apply_blank_original_fallback
from cashflow_direct.models import MaterialityAmounts

def test_结构未确认不得按低金额强行补列():
    c=replace(component(-80407),anomalies=('structure_pending',))
    d=decision(candidate_status='structure_pending',candidate_item_ids=(),summary_candidate_item_ids=(),account_path_candidate_item_ids=())
    out=apply_blank_original_fallback(c,d,MaterialityAmounts(25000000,12500000,1000000),ORDERED_ITEMS,DIRECTIONS)
    assert out==d and not out.resolved

def test_零现金分配也不能消除结构待办():
    c=replace(component(0),anomalies=('structure_pending',))
    d=decision(candidate_status='structure_pending')
    assert apply_blank_original_fallback(c,d,MaterialityAmounts(25000000,12500000,1000000),ORDERED_ITEMS,DIRECTIONS)==d

from dataclasses import replace
from tests.test_business_structure import entry,build

def test_有证据的现金组不被无关未决行阻断():
    rows=(replace(entry(1,'应付职工薪酬_社会保险',7000,'职工社保'),counterpart_name='银行存款'),
          replace(entry(2,'其他应收款_代垫款',3000,'往来对象不详'),counterpart_name=''),
          replace(entry(3,'银行存款',-7000,'支付费用'),counterpart_name='应付职工薪酬'),
          replace(entry(4,'银行存款',-3000,'支付费用'),counterpart_name='其他应收款'))
    result=build(rows)
    resolved=[c for c in result.components if 'structure_pending' not in c.anomalies]
    assert sum(c.cash_delta_cent for c in resolved)==-7000
    assert sum(c.cash_delta_cent for c in result.components)==-10000
    assert result.structure_requests and not result.structure_requests[0].candidate_plans
    assert sum(s.amount_cent for s in result.structure_plans[0].pending_cash_shares)==-3000

def test_等额不同来源仍保留待办():
    rows=tuple(replace(entry(i,a,v,s),counterpart_name=c,source_row_hash=str(i)) for i,a,v,s,c in (
        (1,'其他应收款_个人社保',1000,'个人费用','银行存款'),
        (2,'其他应收款_单位往来',1000,'代付费用','银行存款'),
        (3,'银行存款',-1000,'支付费用','其他应收款'),
        (4,'银行存款',-1000,'支付费用','其他应收款')))
    result=build(rows)
    assert result.structure_requests and not result.structure_requests[0].candidate_plans
    assert all('structure_pending' in c.anomalies for c in result.components)

from pathlib import Path
from dataclasses import replace
from tests.test_business_structure import entry, build
from tests.test_structure_pipeline import start_run
from cashflow_direct import pipeline as p

def test_连续两套分录保留全部现金():
    rows=[]
    for offset in (0,4):
        for number,account,amount,summary,counterpart in (
            (1,'管理费用_办公费',3000,'报销纸张','银行存款'),
            (2,'管理费用_办公费',3000,'报销文件夹','银行存款'),
            (3,'银行存款',-3000,'报销办公用品','管理费用'),
            (4,'银行存款',-3000,'报销办公用品','管理费用'),
        ):
            rows.append(replace(entry(number+offset,account,amount,summary),counterpart_name=counterpart,source_row_hash='相同现金原行' if amount<0 else summary))
    result=build(tuple(rows))
    assert not result.structure_requests
    assert sum(c.cash_delta_cent for c in result.components)==-12000
    assert len(result.source_allocations)==4

def test_新增摘要不得绕过语义门禁(tmp_path):
    run=start_run(tmp_path)
    p.run_classification(run)
    s=p._load_state(run)
    s['summary_semantics']['results']=[]
    s['summary_semantics_completed']=True
    s.pop('classification_summary',None)
    p._save_state(run,s)
    p.run_classification(run)
    after=p._load_state(run)
    assert any(r['summary']=='支付银行手续费' for r in after['summary_semantics']['results'])

def test_未提交任务不计技术失败():
    state={}
    p._register_ai_technical_attempts(state,[{'task_id':'已提交'}],[],[],['尚未提交'])
    assert not state['ai_task_attempts']


def test_诊断导出保留已有正式文件及重算记录(tmp_path):
    import hashlib
    from tests.fixture_factory import write_end_to_end_case

    run = p.run_preflight(
        write_end_to_end_case(tmp_path),
        ('1000000', '750000', '50000'), tmp_path,
    ).run_dir
    state = p._load_state(run)
    p.confirm_cash_scope(run, {
        candidate['account_key']: 'include'
        for candidate in state['cash_scope_proposal']['candidates']
    })
    mark_dictionary_complete(run)
    p.run_classification(run)
    _complete_all_pending_ai(run, tmp_path)
    final = p.finalize_run(run)
    before = p._load_state(run)
    assert final.overall_status == '最终可使用'
    digest = hashlib.sha256(final.workbook_path.read_bytes()).hexdigest()
    assert before['excel_recalculation']['workbook_sha256'] == digest

    diagnostic = p.export_diagnostic_run(run)
    after = p._load_state(run)
    assert after.get('excel_recalculation') == before['excel_recalculation']
    assert after['workbook_path'] == before['workbook_path']
    assert after['overall_status'] == before['overall_status']
    assert after['stage'] == before['stage']
    assert hashlib.sha256(final.workbook_path.read_bytes()).hexdigest() == digest
    assert diagnostic.workbook_path != final.workbook_path
    assert p.finalize_run(run).workbook_path == final.workbook_path


def test_状态保存保留中文金额且编码失败不覆盖原记录(tmp_path):
    import codecs
    import json

    state = {'摘要': '支付职工费用', '金额分': -123456789, '结果': [None, True, {'项目': 'CFO-05'}]}
    p._save_state(tmp_path, state)
    path = tmp_path / '计算留痕数据' / '运行状态.json'
    original = path.read_bytes()
    assert original.startswith(codecs.BOM_UTF8)
    assert json.loads(original.decode('utf-8-sig')) == state
    with pytest.raises(TypeError):
        p._save_state(tmp_path, {'不支持的对象': object()})
    assert path.read_bytes() == original
