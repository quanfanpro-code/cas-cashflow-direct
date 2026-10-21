from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cashflow_direct.paired_input import prepare_paired_input
from cashflow_direct.normalization import normalize_dataset
from cashflow_direct.semantic_mapping import DatasetMapping, infer_dataset_mappings
from cashflow_direct.workbook_structure import scan_workbook


def _book(path: Path, title: str, headers: list[str], rows: list[list[object]]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()


def _files(tmp_path: Path) -> tuple[Path, Path, Path]:
    flow = tmp_path / "现流.xlsx"
    journal = tmp_path / "序时账.xlsx"
    output = tmp_path / "配对.xlsx"
    _book(flow, "明细", ["2026年月", "2026年日", "凭证号", "现金流量项目", "摘要", "方向", "金额"], [
        [1, 4, "记-0006", "销售商品收到的现金", "收回货款", "流入", 100],
        [1, 4, "记-0006", "支付给职工的现金", "支付借款", "流出", 30],
        [1, 4, "记-0007", "销售商品收到的现金", "同摘要", "流入", 20],
        [1, 4, "记-0007", "销售商品收到的现金", "同摘要", "流入", 20],
    ])
    _book(journal, "原始账", ["会计月", "凭证编号", "凭证种类", "业务说明", "上级科目名称", "一级科目", "科目编号", "科目名称", "借方发生额", "贷方发生额"], [
        ["2026-01-04", 6, "记", "收回货款", "\\银行存款", "银行存款", 100201, "一般户", 100, 0],
        ["2026-01-04", 6, "记", "收回货款", "\\主营业务收入", "主营业务收入", 6001, "销售", 0, 100],
        ["2026-01-04", 6, "记", "支付借款", "\\银行存款", "银行存款", 100201, "一般户", 0, 30],
        ["2026-01-04", 6, "记", "支付借款", "\\其他应收款", "其他应收款", 1221, "员工", 30, 0],
        ["2026-01-04", 7, "记", "同摘要", "\\银行存款", "银行存款", 100201, "一般户", 20, 0],
        ["2026-01-04", 7, "记", "同摘要", "\\银行存款", "银行存款", 100201, "一般户", 20, 0],
        ["2026-01-04", 7, "记", "同摘要", "\\主营业务收入", "主营业务收入", 6001, "销售", 0, 40],
        ["2026-01-05", 8, "记", "转存", "\\银行存款", "银行存款", 100201, "一般户", 10, 0],
        ["2026-10-01", 9, "记", "期间外", "\\银行存款", "银行存款", 100201, "一般户", 1, 0],
    ])
    return flow, journal, output


def test_prepare_preserves_original_voucher_and_grouped_duplicates(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    result = prepare_paired_input(flow, journal, output)
    assert (result.journal_rows, result.matched_cash_rows, result.duplicate_groups) == (8, 4, 1)
    book = load_workbook(output, read_only=True, data_only=True)
    rows = list(book.active.values)
    book.close()
    assert rows[1][5] == "银行存款_一般户"
    assert rows[1][8] == "销售商品收到的现金"
    assert rows[2][8] is None
    assert rows[3][8] == "支付给职工的现金"
    assert rows[5][10] == "4,5"
    assert rows[8][8] is None  # 没有现流标签的原始现金分录仍在
    audit = json.loads(result.report.read_text(encoding="utf-8"))
    assert audit["unlabeled_cash_rows"] == 1
    assert audit["journal_sheet"] == "原始账"
    mappings = infer_dataset_mappings(scan_workbook(output))
    assert len(mappings) == 1 and isinstance(mappings[0], DatasetMapping)
    normalized = normalize_dataset(output, "synthetic", mappings[0])
    assert len(normalized.entries) == 8 and not normalized.errors
    assert sum(bool(entry.original_flow_item) for entry in normalized.entries) == 4


def test_ambiguous_project_in_same_identical_group_stops(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(flow)
    book.active[5][3].value = "其他项目"
    book.save(flow)
    book.close()
    with pytest.raises(ValueError, match="多个原项目"):
        prepare_paired_input(flow, journal, output)
    assert not output.exists()


def test_missing_original_cash_leg_stops(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(flow)
    book.active[2][6].value = 99
    book.save(flow)
    book.close()
    with pytest.raises(ValueError, match="无法核对"):
        prepare_paired_input(flow, journal, output)
    assert not output.exists()


def test_negative_flow_reversal_uses_effective_cash_direction(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(flow)
    book.active[2][5].value = "流出"
    book.active[2][6].value = -100
    book.save(flow)
    book.close()
    result = prepare_paired_input(flow, journal, output)
    assert result.matched_cash_rows == 4


def test_derived_counterpart_sheet_is_not_selected(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(journal)
    derived = book.create_sheet("生成结果", 0)
    derived.append(["会计月", "凭证编号", "凭证种类", "业务说明", "科目名称", "科目编号", "借方发生额", "贷方发生额", "对方科目"])
    derived.append(["2026-01-04", 6, "记", "收回货款", "银行存款", 100201, 100, 0, "主营业务收入"])
    book.save(journal)
    book.close()
    result = prepare_paired_input(flow, journal, output)
    assert json.loads(result.report.read_text(encoding="utf-8"))["journal_sheet"] == "原始账"


def test_surplus_identical_cash_leg_stops(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(journal)
    book.active.append(["2026-01-04", 6, "记", "收回货款", "\\银行存款", "银行存款", 100201, "一般户", 100, 0])
    book.save(journal)
    book.close()
    with pytest.raises(ValueError, match="无法核对|金额不一致"):
        prepare_paired_input(flow, journal, output)


def test_title_rows_before_headers_are_allowed(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    for path in (flow, journal):
        book = load_workbook(path)
        book.active.insert_rows(1)
        book.active["A1"] = "本期明细"
        book.save(path)
        book.close()
    result = prepare_paired_input(flow, journal, output)
    assert (result.flow_rows, result.journal_rows) == (4, 8)


def test_same_month_voucher_date_difference_keeps_original_date(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(flow)
    for number in (2, 3):
        book.active[number][1].value = 5
    book.save(flow)
    book.close()
    result = prepare_paired_input(flow, journal, output)
    audit = json.loads(result.report.read_text(encoding="utf-8"))
    assert audit["month_matched_flow_rows"] == 2
    book = load_workbook(output, read_only=True, data_only=True)
    assert book.active["A2"].value == "2026-01-04"
    book.close()


def test_voucher_net_fallback_carries_unique_project_per_direction(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(flow)
    book.active.append([1, 6, "记-0010", "销售商品收到的现金", "平台结算", "流入", 97])
    book.active.append([1, 6, "记-0010", "支付其他经营现金", "平台结算", "流出", 2])
    book.save(flow)
    book.close()
    book = load_workbook(journal)
    sheet = book.active
    sheet.append(["2026-01-06", 10, "记", "平台结算", "\\其他货币资金", "其他货币资金", 101501, "平台款", 105, 0])
    sheet.append(["2026-01-06", 10, "记", "平台结算", "\\其他货币资金", "其他货币资金", 101501, "平台款", 0, 10])
    sheet.append(["2026-01-06", 10, "记", "平台结算", "\\应收账款", "应收账款", 1122, "客户", 0, 105])
    sheet.append(["2026-01-06", 10, "记", "平台结算", "\\其他应收款", "其他应收款", 1221, "平台", 10, 0])
    book.save(journal)
    book.close()
    result = prepare_paired_input(flow, journal, output)
    audit = json.loads(result.report.read_text(encoding="utf-8"))
    assert audit["net_matched_vouchers"] == 1
    book = load_workbook(output, read_only=True, data_only=True)
    rows = list(book.active.values)
    book.close()
    assert rows[-4][8] == "销售商品收到的现金"
    assert rows[-3][8] == "支付其他经营现金"
    assert rows[-4][11] == "凭证净额"


def test_net_fallback_stops_if_one_direction_has_two_projects(tmp_path: Path) -> None:
    flow, journal, output = _files(tmp_path)
    book = load_workbook(flow)
    book.active.append([1, 6, "记-0010", "项目甲", "同一结算", "流入", 60])
    book.active.append([1, 6, "记-0010", "项目乙", "同一结算", "流入", 40])
    book.save(flow)
    book.close()
    book = load_workbook(journal)
    book.active.append(["2026-01-06", 10, "记", "同一结算", "\\银行存款", "银行存款", 100201, "一般户", 100, 0])
    book.save(journal)
    book.close()
    with pytest.raises(ValueError, match="同方向.*多个原项目"):
        prepare_paired_input(flow, journal, output)
